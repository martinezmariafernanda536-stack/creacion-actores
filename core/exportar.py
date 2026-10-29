"""Genera el Excel con el formato oficial de la plantilla de creación de terceros HERMES."""
from __future__ import annotations

from copy import copy
from datetime import datetime
from io import BytesIO

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, Side

from .config import PLANTILLA_BASE

# Columna de la plantilla -> campo del sistema
MAPEO = [
    ("A", "solicitante"), ("B", "fecha_solicitud"), ("C", "aplicativo"), ("D", "estado"),
    ("E", "tipo_organizacion"), ("F", "tipo_entidad"), ("G", "tipo_identificacion"),
    ("H", "identificacion"), ("I", "nombre_organizacion"), ("J", "direccion"),
    ("K", "continente"), ("L", "pais"), ("M", "departamento"), ("N", "municipio"),
    ("O", "email"), ("P", "telefono_fijo"), ("Q", "telefono_movil"),
]
FILA_INICIO = 8


def generar_excel(df: pd.DataFrame) -> bytes:
    wb = openpyxl.load_workbook(PLANTILLA_BASE)
    ws = wb["Organización"]
    estilo = ws.cell(FILA_INICIO, 1)
    borde = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
    for i, (_, r) in enumerate(df.iterrows()):
        fila = FILA_INICIO + i
        for col, campo in MAPEO:
            c = ws[f"{col}{fila}"]
            val = r.get(campo, "")
            if campo == "fecha_solicitud" and val:
                try:
                    val = datetime.strptime(str(val)[:10], "%Y-%m-%d")
                    c.number_format = "DD/MM/YYYY"
                except ValueError:
                    pass
            if campo in ("identificacion", "telefono_fijo", "telefono_movil"):
                c.number_format = "@"  # texto: conserva ceros y dígitos completos
            c.value = val if val != "" else None
            c.font = copy(estilo.font) if estilo.has_style else Font(name="Calibri", size=10)
            c.border = borde
            c.alignment = Alignment(vertical="center", wrap_text=campo in ("nombre_organizacion", "direccion"))

    # Hoja de control con los radicados del sistema (no altera la hoja oficial)
    ctrl = wb.create_sheet("Control sistema")
    campos = ["radicado", "lote", "estado", "existe_en_hermes", "numero_itop", "codigo_hermes",
              "codigo_ciclope", "fuente_dato", "observaciones_solicitante"]
    ctrl.append(campos)
    for c in ctrl[1]:
        c.font = Font(bold=True)
    for _, r in df.iterrows():
        ctrl.append([str(r.get(k, "")) for k in campos])
    for col in "ABCDEFGHI":
        ctrl.column_dimensions[col].width = 22

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
